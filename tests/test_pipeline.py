"""End-to-end pipeline run on the recorded Kortrijk data with every external call mocked."""

from __future__ import annotations

import csv
import json
import re
from urllib.parse import parse_qs, urlsplit

import pytest
import responses
from openpyxl import load_workbook

from conftest import fixture_json, fixture_text
from dataharvest.enrich.search import DDG_HTML_URL
from dataharvest.enrich.vies import VIES_URL
from dataharvest.models import FieldStatus, RecordStatus
from dataharvest.pipeline import Pipeline
from dataharvest.processing.verify import MailDomainChecker, MxResult
from dataharvest.sources.nominatim import NOMINATIM_URL
from dataharvest.sources.osm_overpass import OVERPASS_ENDPOINTS


def site(name: str, *, email: str = "", phone: str = "", vat: str = "", extra: str = "") -> str:
    return f"""<!DOCTYPE html><html><head><title>{name} - officiële website</title>
    <meta name="description" content="{name} in Kortrijk."></head>
    <body><h1>Welkom bij {name}</h1>
    {'<p>Mail: <a href="mailto:' + email + '">' + email + '</a></p>' if email else ''}
    {'<p>Tel: <a href="tel:' + phone + '">' + phone + '</a></p>' if phone else ''}
    {'<footer>BTW ' + vat + '</footer>' if vat else ''}
    {extra}
    <a href="/contact">Contact</a></body></html>"""


def ddg_page(url: str, title: str, snippet: str) -> str:
    from urllib.parse import quote

    return f"""<html><body><div id="links"><div class="result web-result"><h2 class="result__title">
    <a class="result__a" href="//duckduckgo.com/l/?uddg={quote(url, safe='')}&amp;rut=abc">{title}</a></h2>
    <a class="result__snippet" href="#">{snippet}</a></div></div></body></html>"""


def register_network() -> None:
    responses.add(responses.GET, NOMINATIM_URL, json=fixture_json("nominatim_kortrijk.json"))
    responses.add(responses.POST, OVERPASS_ENDPOINTS[0], json=fixture_json("overpass_kortrijk.json"))
    # company websites
    responses.add(responses.GET, "https://hollywok.be/", body=site("Hollywok", email="info@hollywok.be", phone="+32 56 20 55 88", vat="BE 0473.191.041",
                                                                  extra='<a href="https://www.facebook.com/HollywokKortrijk/">fb</a>'), content_type="text/html")
    responses.add(responses.GET, "https://hollywok.be/contact", body=site("Hollywok", email="info@hollywok.be"), content_type="text/html")
    responses.add(responses.GET, "http://www.y-notthai.com/", body=site("Y-Not Thai", email="info@y-notthai.com", phone="+32 56 51 20 48"), content_type="text/html")
    responses.add(responses.GET, "https://kostbar-marke.be/", body=site("Kostbar", email="els@kostbar-marke.be"), content_type="text/html")
    responses.add(responses.GET, "https://www.frituurnatuur.be/", body=site("Frituur Natuur", vat="BE0629.985.405"), content_type="text/html")
    responses.add(responses.GET, "https://alvaro-koffie.be/", body=site("Alvaro", email="hallo@alvaro-koffie.be"), content_type="text/html")
    responses.add(responses.GET, "https://www.bastakortrijk.be/", status=500)

    def ddg(request):
        query = parse_qs(urlsplit(request.url).query).get("q", [""])[0]
        if "Alvaro" in query:
            return 200, {"Content-Type": "text/html"}, ddg_page("https://alvaro-koffie.be/", "Alvaro koffiebar Heule", "Alvaro - koffie in Heule, Kortrijk")
        return 200, {"Content-Type": "text/html"}, fixture_text("ddg_results.html")

    responses.add_callback(responses.GET, DDG_HTML_URL, callback=ddg)
    responses.add(responses.GET, VIES_URL.format(country="BE", number="0629985405"), json=fixture_json("vies_valid.json"))
    hollywok_vies = dict(fixture_json("vies_valid.json"), name="BV HOLLYWOK", address="President Kennedylaan 100\n8500 Kortrijk", vatNumber="0473191041")
    responses.add(responses.GET, VIES_URL.format(country="BE", number="0473191041"), json=hollywok_vies)
    # everything else (all other company websites) is a 404 - but never shadow the specific mocks above
    catch_all = re.compile(r"^(?!.*(hollywok\.be|y-notthai|kostbar-marke|frituurnatuur|alvaro-koffie|duckduckgo|ec\.europa|nominatim|overpass)).*")
    responses.add(responses.GET, catch_all, status=404, body="<html><title>Not found</title></html>", content_type="text/html")


@pytest.fixture
def fake_dns(monkeypatch):
    def resolve(self, domain: str) -> MxResult:
        if "does-not-exist" in domain:
            return MxResult(domain, False, "domain does not exist (NXDOMAIN)")
        return MxResult(domain, True, "MX: mail.example")

    monkeypatch.setattr(MailDomainChecker, "_resolve", resolve)


@responses.activate
def test_full_run(leads_project, http, fake_dns):
    register_network()
    result = Pipeline(leads_project, http=http).run(export=True)
    records = result.records
    by_name = {r.get("company_name"): r for r in records if r.get("company_name")}

    # --- extraction, normalisation, ids
    assert len(records) == 60
    assert all(re.match(r"TES-\d{4}", r.record_id) for r in records)
    hollywok = by_name["Hollywok"]
    assert hollywok.get("phone") == "+3256205588" and hollywok.get("website") == "https://hollywok.be/"

    # --- website enrichment confirms contact details and finds new ones
    assert hollywok.field_status("website") == FieldStatus.VERIFIED
    assert hollywok.field_status("email") == FieldStatus.VERIFIED and "confirmed" in hollywok.fields["email"].note
    assert hollywok.field_status("phone") == FieldStatus.VERIFIED
    assert hollywok.get("facebook") == "https://www.facebook.com/HollywokKortrijk"
    assert hollywok.get("vat_number") == "BE0473191041" and hollywok.field_status("vat_number") == FieldStatus.VERIFIED
    assert hollywok.get("legal_name") == "BV HOLLYWOK"
    assert hollywok.status == RecordStatus.VERIFIED, hollywok.flags

    # --- conflict between source and website is flagged, not silently overwritten
    thai = by_name["Thai Lin"]
    assert thai.get("email") == "info@thailin.be" and thai.field_status("email") == FieldStatus.CONFLICT
    assert "info@y-notthai.com" in thai.fields["email"].candidates
    assert thai.status == RecordStatus.NEEDS_REVIEW

    # --- dead / unreachable websites
    basta = by_name["Basta!"]
    assert basta.field_status("website") == FieldStatus.INVALID and any("unreachable" in f for f in basta.flags)
    assert basta.status == RecordStatus.NEEDS_REVIEW

    # --- duplicates merged (same e-mail + website, similar names)
    kostbar, kostbaar = by_name["Kostbar"], by_name["Kostba(a)r"]
    assert kostbar.duplicate_group == kostbaar.duplicate_group and kostbaar.duplicate_of == kostbar.record_id
    assert kostbaar.status == RecordStatus.EXCLUDED
    assert any(g.merged and kostbar.record_id in g.member_ids for g in result.groups)

    # --- website found through web search, then confirmed on the page
    alvaro = by_name["Alvaro"]
    assert alvaro.get("website") == "https://alvaro-koffie.be/"
    assert alvaro.fields["website"].source == "search:duckduckgo"
    assert alvaro.field_status("website") == FieldStatus.VERIFIED
    assert not any("found via web search - verify" in f for f in alvaro.flags)
    assert alvaro.get("email") == "hallo@alvaro-koffie.be"

    # --- VIES on a VAT number that came from the source
    frituur = by_name["Frituur Natuur"]
    assert frituur.field_status("vat_number") == FieldStatus.VERIFIED and frituur.get("legal_name") == "BV HUMICO"
    assert any("legal entity name" in f for f in frituur.flags)

    # --- unusable records are excluded, not silently dropped
    excluded = [r for r in records if r.status == RecordStatus.EXCLUDED and not r.duplicate_of]
    assert len(excluded) == 4 and all("missing company name" in " ".join(r.flags) for r in excluded)

    # --- two entries with the same name and website but different addresses: merged, address conflict flagged
    btc = [r for r in records if r.get("company_name") == "Bistro Tout Court"]
    assert len(btc) == 2 and sum(1 for r in btc if r.duplicate_of) == 1
    master = next(r for r in btc if not r.duplicate_of)
    assert master.status == RecordStatus.NEEDS_REVIEW and any("disagree on" in f and "street" in f for f in master.flags)

    # --- report and outputs
    rep = result.report
    assert rep.records_total == 60 and rep.records_delivered == 60 - 4 - 2
    assert rep.status_counts["EXCLUDED"] == 6
    assert set(result.outputs) >= {"xlsx", "csv", "json", "report_md", "report_json"}
    for p in result.outputs.values():
        assert p.exists()

    wb = load_workbook(result.outputs["xlsx"])
    assert wb.sheetnames == ["Summary", "Leads", "Needs Review", "Duplicates", "Run Log", "Data Dictionary"]
    ws = wb["Leads"]
    assert ws.max_row - 1 == rep.records_delivered
    headers = [c.value for c in ws[1]]
    assert headers[:3] == ["Record ID", "Company name", "Legal entity name"] and "Verification status" in headers
    assert ws.freeze_panes == "C2" and ws.auto_filter.ref.startswith("A1:")
    review = wb["Needs Review"]
    assert review.max_row - 1 == rep.status_counts["NEEDS_REVIEW"] + 4  # + excluded (nameless) records
    dups = wb["Duplicates"]
    assert dups.cell(row=2, column=1).value.startswith("DUP-")
    summary = wb["Summary"]
    assert any(isinstance(c.value, str) and c.value.startswith("=COUNTA(") for row in summary.iter_rows() for c in row)

    with open(result.outputs["csv"], encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    assert len(rows) - 1 == rep.records_delivered and rows[0][1] == "Company name"

    payload = json.loads(result.outputs["json"].read_text(encoding="utf-8"))
    assert len(payload["records"]) == 60 and payload["records"][0]["fields"]["company_name"]["source"]
    md = result.outputs["report_md"].read_text(encoding="utf-8")
    assert "# Run report" in md and "VERIFIED" in md


@responses.activate
def test_offline_run_skips_network(leads_project, http):
    responses.add(responses.GET, NOMINATIM_URL, json=fixture_json("nominatim_kortrijk.json"))
    responses.add(responses.POST, OVERPASS_ENDPOINTS[0], json=fixture_json("overpass_kortrijk.json"))
    result = Pipeline(leads_project, http=http, offline=True, limit=10).run(export=False)
    assert len(result.records) == 10 and len(responses.calls) == 2
    assert all(r.field_status("website") in (FieldStatus.UNVERIFIED, FieldStatus.MISSING) for r in result.records)
    assert not result.outputs
