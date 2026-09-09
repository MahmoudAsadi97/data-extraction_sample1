"""Regression tests for real-world failure modes: blocked sites, transient API errors, locked files, odd data."""

from __future__ import annotations

import json
from pathlib import Path

import requests
import responses
import yaml
from openpyxl import load_workbook

from conftest import fixture_json
from dataharvest.config import SourceConfig, load_project
from dataharvest.enrich.vies import VIES_URL, ViesClient
from dataharvest.enrich.website import SiteExtraction, WebsiteEnricher
from dataharvest.export import export_excel
from dataharvest.http import cacheable
from dataharvest.models import FieldStatus, Record, RecordStatus
from dataharvest.pipeline import Pipeline
from dataharvest.processing.verify import MailDomainChecker
from dataharvest.report import RunReport
from dataharvest.schema import load_schema
from dataharvest.sources.nominatim import NOMINATIM_URL
from dataharvest.sources.osm_overpass import OVERPASS_ENDPOINTS, OsmOverpassSource

LEADS = load_schema("leads")


class TestBlockedAndFallbacks:
    @responses.activate
    def test_robots_blocked_site_is_not_reported_dead(self, http, leads_project):
        responses.add(responses.GET, "https://blocked.example.org/robots.txt", body="User-agent: *\nDisallow: /\n")
        http.respect_robots = True
        site = WebsiteEnricher(http, respect_robots=True).analyse("https://blocked.example.org/", "Blocked")
        assert not site.ok and site.blocked and site.liveness == "blocked"
        rec = Record(record_id="X", source="osm")
        rec.set("company_name", "Blocked", "osm")
        rec.set("website", "https://blocked.example.org/", "osm")
        Pipeline(leads_project, http=http)._apply_site_extraction(rec, site)
        assert rec.field_status("website") == FieldStatus.UNVERIFIED
        assert not any("unreachable" in f for f in rec.flags) and any(f.startswith("note:") for f in rec.flags)

    @responses.activate
    def test_bot_challenge_is_blocked_not_dead(self, http):
        responses.add(responses.GET, "https://cf.example.org/", status=403, body="<html><title>Just a moment...</title></html>",
                      headers={"Server": "cloudflare"}, content_type="text/html")
        site = WebsiteEnricher(http, respect_robots=False).analyse("https://cf.example.org/")
        assert site.liveness == "blocked"
        responses.add(responses.GET, "https://gone.example.org/", status=404, body="<html>no</html>", content_type="text/html")
        assert WebsiteEnricher(http, respect_robots=False).analyse("https://gone.example.org/").liveness == "dead"

    @responses.activate
    def test_https_failure_falls_back_to_http(self, http):
        responses.add(responses.GET, "https://oldsite.example.org/", body=requests.exceptions.SSLError("handshake"))
        responses.add(responses.GET, "http://oldsite.example.org/", body="<html><title>Old Site</title><body>Old Site bv</body></html>", content_type="text/html")
        site = WebsiteEnricher(http, respect_robots=False).analyse("https://oldsite.example.org/", "Old Site")
        assert site.ok and site.liveness == "live" and site.name_match and "plain http" in site.note

    def test_email_case_does_not_create_conflicts(self, http, leads_project):
        rec = Record(record_id="X", source="csv")
        rec.set("company_name", "Thai Lin", "csv")
        rec.set("email", "info@thailin.be", "csv")  # normalised form is lower-case
        rec.set("website", "https://thailin.be/", "csv")
        site = SiteExtraction(url="https://thailin.be/", ok=True, status=200, final_url="https://thailin.be/", emails=["info@thailin.be"], name_match=True)
        Pipeline(leads_project, http=http)._apply_site_extraction(rec, site)
        assert rec.field_status("email") == FieldStatus.VERIFIED


class TestTransientErrorsAreNotCached:
    def test_cache_filter_rejects_overpass_and_vies_errors(self):
        def fake(url: str, payload: dict) -> requests.Response:
            r = requests.Response()
            r.url = url
            r.status_code = 200
            r._content = json.dumps(payload).encode()
            return r

        assert cacheable(fake("https://overpass-api.de/api/interpreter", {"elements": [], "remark": "runtime error: Query timed out"})) is False
        assert cacheable(fake("https://overpass-api.de/api/interpreter", {"elements": [{"id": 1}]})) is True
        assert cacheable(fake(VIES_URL.format(country="BE", number="1"), fixture_json("vies_busy.json"))) is False
        assert cacheable(fake(VIES_URL.format(country="BE", number="1"), fixture_json("vies_valid.json"))) is True
        assert cacheable(fake("https://example.org/page", {"anything": 1})) is True

    @responses.activate
    def test_overpass_runtime_error_tries_next_mirror(self, http, tmp_path):
        from test_sources import project

        responses.add(responses.GET, NOMINATIM_URL, json=fixture_json("nominatim_kortrijk.json"))
        responses.add(responses.POST, OVERPASS_ENDPOINTS[0], json={"elements": [], "remark": "runtime error: Query timed out in \"query\""})
        responses.add(responses.POST, OVERPASS_ENDPOINTS[1], json=fixture_json("overpass_kortrijk.json"))
        src = OsmOverpassSource(SourceConfig(type="osm_overpass", area="Kortrijk", tags=["amenity=cafe"]), http, project(tmp_path))
        assert len(list(src.extract())) == 60

    @responses.activate
    def test_vies_retries_bypass_cache_then_succeed(self, http):
        url = VIES_URL.format(country="BE", number="0629985405")
        responses.add(responses.GET, url, json=fixture_json("vies_busy.json"))
        responses.add(responses.GET, url, json=fixture_json("vies_valid.json"))
        result = ViesClient(http, delay=0).check("BE0629985405")
        assert result.valid is True and len(responses.calls) == 2


class TestDnsAndExcelRobustness:
    def test_nonameservers_is_unknown_not_invalid(self, monkeypatch):
        import dns.resolver

        class Boom:
            lifetime = 0

            def resolve(self, *a, **k):
                raise dns.resolver.NoNameservers()

        monkeypatch.setattr(dns.resolver, "Resolver", lambda: Boom())
        result = MailDomainChecker(http=None).check("example.org")
        assert result.deliverable is None

    def test_excel_survives_control_characters_and_formula_like_text(self, tmp_path):
        rec = Record(record_id="X-0001", source="osm")
        rec.set("company_name", "=SUM(A1)\x0b weird", "osm")
        rec.set("category", "cafe", "osm")
        rec.set("description", "meta \x0c description\x1f", "website")
        rec.status = RecordStatus.UNVERIFIED
        report = RunReport("p", "P")
        report.warn("warning with \x0b control char")
        report.summarise([rec], [], LEADS)
        path = export_excel(tmp_path / "x.xlsx", [rec], [], report, LEADS, {"name": "p", "title": "P"})
        ws = load_workbook(path)["Leads"]
        assert ws.cell(row=2, column=2).value == "=SUM(A1) weird" and ws.cell(row=2, column=2).data_type == "s"

    def test_locked_output_falls_back_to_another_name(self, tmp_path, monkeypatch, leads_project):
        from dataharvest.export import _write

        pipeline = Pipeline(leads_project, offline=True)
        calls = []

        def writer(path: Path) -> Path:
            calls.append(path)
            if len(calls) == 1:
                raise PermissionError("locked")
            path.write_text("ok", encoding="utf-8")
            return path

        out = _write(pipeline, "xlsx", tmp_path / "result.xlsx", writer)
        assert out is not None and out.exists() and out.name != "result.xlsx"
        assert any("locked" in w for w in pipeline.report.warnings)


def test_project_deadline_as_yaml_date(tmp_path):
    data = {"project": {"name": "p", "deadline": "2026-10-01"}, "schema": "leads", "sources": [{"type": "csv_import", "path": "x.csv"}]}
    text = yaml.safe_dump(data).replace("'2026-10-01'", "2026-10-01")  # unquoted -> YAML date
    path = tmp_path / "p.yaml"
    path.write_text(text, encoding="utf-8")
    assert load_project(path).project.deadline == "2026-10-01"


class TestDnsFailures:
    @responses.activate
    def test_temporary_resolver_failure_is_unchecked_not_dead(self, http, leads_project):
        from urllib3.exceptions import NameResolutionError

        err = requests.exceptions.ConnectionError(NameResolutionError("flaky.example.org", None, "[Errno -3] Temporary failure in name resolution"))
        responses.add(responses.GET, "https://flaky.example.org/", body=err)
        site = WebsiteEnricher(http, respect_robots=False).analyse("https://flaky.example.org/", "Flaky")
        assert site.dns_failure and site.dns_temporary and site.liveness == "unchecked"
        assert len(responses.calls) == 1  # no pointless http:// retry when the name itself does not resolve
        rec = Record(record_id="X", source="osm")
        rec.set("company_name", "Flaky", "osm")
        rec.set("website", "https://flaky.example.org/", "osm")
        Pipeline(leads_project, http=http)._apply_site_extraction(rec, site)
        assert rec.field_status("website") == FieldStatus.UNVERIFIED and any("name resolution failed" in f for f in rec.flags)
        assert not any("unreachable" in f for f in rec.flags)

    @responses.activate
    def test_unknown_host_is_unreachable_when_dns_works(self, http, leads_project):
        from urllib3.exceptions import NameResolutionError

        err = requests.exceptions.ConnectionError(NameResolutionError("gone.example.org", None, "[Errno -2] Name or service not known"))
        responses.add(responses.GET, "https://gone.example.org/", body=err)
        site = WebsiteEnricher(http, respect_robots=False).analyse("https://gone.example.org/", "Gone")
        assert site.dns_failure and not site.dns_temporary and site.liveness == "unreachable"
        rec = Record(record_id="X", source="osm")
        rec.set("company_name", "Gone", "osm")
        rec.set("website", "https://gone.example.org/", "osm")
        pipeline = Pipeline(leads_project, http=http)
        pipeline.dns_broken = False
        pipeline._apply_site_extraction(rec, site)
        assert rec.field_status("website") == FieldStatus.INVALID and any("unreachable" in f for f in rec.flags)
        pipeline.dns_broken = True  # ... but when the machine's DNS is known to be broken, nothing is declared dead
        rec2 = Record(record_id="Y", source="osm")
        rec2.set("company_name", "Gone", "osm")
        rec2.set("website", "https://gone.example.org/", "osm")
        pipeline._apply_site_extraction(rec2, site)
        assert rec2.field_status("website") == FieldStatus.UNVERIFIED

    def test_dns_health_check(self, monkeypatch):
        import socket

        from dataharvest.http import dns_healthy

        monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **k: (_ for _ in ()).throw(socket.gaierror("boom")))
        assert dns_healthy() is False
        monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **k: [("ok",)])
        assert dns_healthy() is True
