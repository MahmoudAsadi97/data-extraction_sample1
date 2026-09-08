"""Exports, the stand-alone quality audit and the command-line interface."""

from __future__ import annotations

import csv
import json
import shutil

import responses
from openpyxl import load_workbook
from typer.testing import CliRunner

from conftest import FIXTURES, fixture_json
from dataharvest.cli import app
from dataharvest.export import export_csv, export_excel, export_json
from dataharvest.export.columns import build_columns, record_row
from dataharvest.export.gsheets import GoogleSheetsUnavailable, export_google_sheets
from dataharvest.models import DuplicateGroup, FieldStatus, Record, RecordStatus
from dataharvest.qa import audit_file, map_columns
from dataharvest.report import RunReport
from dataharvest.schema import load_schema

LEADS = load_schema("leads")
runner = CliRunner()


def sample_records() -> list[Record]:
    a = Record(record_id="X-0001", source="osm", source_url="https://www.openstreetmap.org/node/1")
    a.set("company_name", "Hollywok", "osm", FieldStatus.VERIFIED, "confirmed on company website")
    a.set("category", "restaurant", "osm")
    a.set("phone", "+3256205588", "osm", FieldStatus.VERIFIED)
    a.set("email", "info@hollywok.be", "website", FieldStatus.VERIFIED, "found on company website")
    a.set("website", "https://hollywok.be/", "osm", FieldStatus.VERIFIED, "live (HTTP 200); company name found on page")
    a.status, a.completeness = RecordStatus.VERIFIED, 40.0
    b = Record(record_id="X-0002", source="osm")
    b.set("company_name", "Thai Lin", "osm")
    b.set("category", "restaurant", "osm")
    b.set("email", "info@thailin.be", "osm", FieldStatus.CONFLICT, "website lists info@y-notthai.com instead")
    b.fields["email"].candidates = ["info@y-notthai.com"]
    b.flag("e-mail differs from the company website (info@y-notthai.com) - review")
    b.status, b.completeness = RecordStatus.NEEDS_REVIEW, 20.0
    c = Record(record_id="X-0003", source="osm")
    c.set("category", "restaurant", "osm")
    c.flag("missing company name - record unusable")
    c.status = RecordStatus.EXCLUDED
    return [a, b, c]


class TestColumnsAndFlatExports:
    def test_columns_and_row(self):
        cols = build_columns(LEADS)
        keys = [c.key for c in cols]
        assert keys[0] == "record_id" and "check:email" in keys and keys[-1] == "reviewer_notes"
        row = record_row(sample_records()[0], cols)
        assert row[0] == "X-0001" and row[keys.index("phone")] == "+32 56 20 55 88"
        assert row[keys.index("check:email")].startswith("verified - found on company website")
        conflict_row = record_row(sample_records()[1], cols)
        assert "other values seen: info@y-notthai.com" in conflict_row[keys.index("check:email")]

    def test_csv_and_json(self, tmp_path):
        records = sample_records()
        report = RunReport("p", "P")
        report.summarise(records, [], LEADS)
        csv_path = export_csv(tmp_path / "out.csv", records, LEADS)
        with csv_path.open(encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.DictReader(fh))
        assert len(rows) == 2 and rows[0]["Company name"] == "Hollywok" and rows[1]["Verification status"] == "NEEDS_REVIEW"
        json_path = export_json(tmp_path / "out.json", records, [], report, LEADS)
        data = json.loads(json_path.read_text(encoding="utf-8"))
        assert len(data["records"]) == 3 and data["records"][1]["fields"]["email"]["candidates"] == ["info@y-notthai.com"]
        assert data["report"]["records_delivered"] == 2


class TestExcel:
    def test_workbook_structure(self, tmp_path):
        records = sample_records()
        groups = [DuplicateGroup("DUP-001", "X-0001", ["X-0001", "X-0002"], "same phone", 95.0, merged=False)]
        report = RunReport("p", "Project P")
        report.summarise(records, groups, LEADS)
        path = export_excel(tmp_path / "out.xlsx", records, groups, report, LEADS, {"name": "p", "title": "Project P", "country": "BE"})
        wb = load_workbook(path)
        assert wb.sheetnames == ["Summary", "Leads", "Needs Review", "Duplicates", "Run Log", "Data Dictionary"]
        ws = wb["Leads"]
        assert ws.max_row == 3  # header + 2 delivered (excluded record left out)
        headers = [c.value for c in ws[1]]
        email_col = headers.index("E-mail") + 1
        assert ws.cell(row=2, column=email_col).hyperlink.target == "mailto:info@hollywok.be"
        website_col = headers.index("Website") + 1
        assert ws.cell(row=2, column=website_col).hyperlink.target == "https://hollywok.be/"
        assert len(ws.conditional_formatting) >= 5
        assert ws.data_validations.dataValidation[0].formula1 == '"Approved,Rejected,Needs fix,Contacted"'
        review = wb["Needs Review"]
        assert review.cell(row=2, column=1).value == "X-0002" and review.cell(row=3, column=1).value == "X-0003"
        dups = wb["Duplicates"]
        assert dups.cell(row=2, column=2).value.startswith("POSSIBLE")
        dictionary = wb["Data Dictionary"]
        assert dictionary.cell(row=2, column=1).value == "Record ID"
        log = wb["Run Log"]
        assert log["A1"].value == "Run log"

    def test_google_sheets_unavailable_without_credentials(self, monkeypatch):
        monkeypatch.delenv("GOOGLE_SERVICE_ACCOUNT_JSON", raising=False)
        try:
            export_google_sheets(sample_records(), [], LEADS, title="x")
        except GoogleSheetsUnavailable as exc:
            assert "GOOGLE_SERVICE_ACCOUNT_JSON" in str(exc) or "gspread" in str(exc)
        else:  # pragma: no cover
            raise AssertionError("expected GoogleSheetsUnavailable")


class TestAudit:
    def test_map_columns(self):
        mapping, ignored = map_columns(["Company", "Website", "Tel", "E-mail", "City", "Postal code", "Weird"], LEADS, {"Tel": "phone"})
        assert mapping == {"Company": "company_name", "Website": "website", "Tel": "phone", "E-mail": "email", "City": "city", "Postal code": "postcode"}
        assert ignored == ["Weird"]

    def test_audit_seed_list(self, tmp_path):
        path = tmp_path / "seed.csv"
        shutil.copy(FIXTURES / "seed_list.csv", path)
        result = audit_file(path, "leads", country="BE")
        assert len(result.records) == 5
        by_name = {r.get("company_name"): r for r in result.records}
        broken = by_name["Unknown Corp"]
        joined = " | ".join(broken.flags)
        assert "postcode" in joined and "phone" in joined and "e-mail" in joined
        assert broken.status == RecordStatus.NEEDS_REVIEW
        assert result.groups and not result.groups[0].merged  # Hollywok twice -> flagged, not merged
        assert by_name["Hollywok"].duplicate_group == by_name["Hollywok Kortrijk"].duplicate_group
        assert result.rows_with_problems >= 2 and result.problems >= 3
        assert by_name["Hollywok"].get("phone") == "+3256205588"  # normalised


class TestCli:
    def test_help_and_info_commands(self):
        assert runner.invoke(app, ["--help"]).exit_code == 0
        assert "dataharvest 1." in runner.invoke(app, ["--version"]).stdout
        out = runner.invoke(app, ["sources"])
        assert out.exit_code == 0 and "osm_overpass" in out.stdout
        out = runner.invoke(app, ["schemas", "leads"])
        assert out.exit_code == 0 and "company_name" in out.stdout

    def test_projects_listing(self, repo_root):
        out = runner.invoke(app, ["projects", "--dir", str(repo_root / "projects")])
        flat = "".join(out.stdout.split())  # rich wraps long cells in narrow terminals
        assert out.exit_code == 0 and "kortrijk_restaurants.yaml" in flat and "invalid" not in flat

    def test_shipped_projects_are_valid(self, repo_root):
        from dataharvest.config import load_project

        files = sorted((repo_root / "projects").glob("*.yaml"))
        assert len(files) >= 5
        for f in files:
            cfg = load_project(f)
            assert cfg.active_sources and cfg.schema_def.fields
            for src in cfg.sources:
                if src.type == "csv_import":
                    assert cfg.resolve_path(src.options()["path"]).exists(), f"{f.name}: input file missing"

    def test_init_creates_valid_project(self, tmp_path):
        out = runner.invoke(app, ["init", "bakeries", "--template", "leads", "--dir", str(tmp_path)])
        assert out.exit_code == 0 and (tmp_path / "bakeries.yaml").exists()
        from dataharvest.config import load_project

        cfg = load_project(tmp_path / "bakeries.yaml")
        assert cfg.project.name == "bakeries" and cfg.schema_def.name == "leads"
        for template in ("companies", "products", "csv"):
            assert runner.invoke(app, ["init", f"p_{template}", "--template", template, "--dir", str(tmp_path)]).exit_code == 0
            load_project(tmp_path / f"p_{template}.yaml")

    def test_validate_optional_and_mapping_flags(self, tmp_path):
        path = tmp_path / "seed.csv"
        shutil.copy(FIXTURES / "seed_list.csv", path)
        out = runner.invoke(app, ["validate", str(path), "--optional", "category", "-m", "Notes=description"])
        assert out.exit_code == 0, out.stdout
        assert "missing required field: category" not in out.stdout

    def test_validate_command(self, tmp_path):
        path = tmp_path / "seed.csv"
        shutil.copy(FIXTURES / "seed_list.csv", path)
        out = runner.invoke(app, ["validate", str(path), "--schema", "leads", "--country", "BE"])
        assert out.exit_code == 0, out.stdout
        assert (tmp_path / "seed_audit.xlsx").exists()
        assert "Unknown Corp" in out.stdout

    @responses.activate
    def test_run_command_offline(self, tmp_path, leads_project):
        from dataharvest.sources.nominatim import NOMINATIM_URL
        from dataharvest.sources.osm_overpass import OVERPASS_ENDPOINTS

        responses.add(responses.GET, NOMINATIM_URL, json=fixture_json("nominatim_kortrijk.json"))
        responses.add(responses.POST, OVERPASS_ENDPOINTS[0], json=fixture_json("overpass_kortrijk.json"))
        out = runner.invoke(app, ["run", str(leads_project.path), "--offline", "--limit", "8", "--out", str(tmp_path / "o")])
        assert out.exit_code == 0, out.stdout
        assert "delivered" in out.stdout and list((tmp_path / "o").glob("*.xlsx"))

    def test_run_missing_project(self):
        out = runner.invoke(app, ["run", "does_not_exist"])
        assert out.exit_code != 0
