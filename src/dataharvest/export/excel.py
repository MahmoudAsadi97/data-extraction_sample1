"""Excel workbook export: formatted data sheet, review queue, duplicates, live summary, run log, data dictionary."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from ..models import DuplicateGroup, Record, RecordStatus
from ..processing.normalize import format_phone_display
from ..report import RunReport
from ..schema import Schema
from .columns import Column, build_columns, record_row, spreadsheet_safe

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
SUBTLE_FONT = Font(italic=True, color="7F7F7F")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL_GREEN = PatternFill("solid", fgColor="C6EFCE")
FILL_YELLOW = PatternFill("solid", fgColor="FFEB9C")
FILL_GREY = PatternFill("solid", fgColor="EDEDED")
FILL_RED = PatternFill("solid", fgColor="FFC7CE")
FILL_BLUE = PatternFill("solid", fgColor="DDEBF7")
FONT_GREEN = Font(color="006100")
FONT_YELLOW = Font(color="9C5700")
FONT_RED = Font(color="9C0006")
FONT_GREY = Font(color="595959")
MAX_ROWS = 100_000

REVIEW_COLUMNS = ("record_id", "status", "flags", "source", "source_url")


def _display(rec: Record, name: str, schema: Schema) -> Any:
    """Cell value for a field, formatted for humans (phones in international format)."""
    value = rec.get(name)
    f = schema.field(name)
    if value and f is not None and f.type == "phone":
        return format_phone_display(str(value))
    return value


def _put(ws: Worksheet, row: int, column: int, value: Any):
    """Write a data cell; text that starts with '=' stays text instead of becoming a formula."""
    cell = ws.cell(row=row, column=column, value=spreadsheet_safe(value))
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    return cell


def _style_header(ws: Worksheet, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 30


def _write_table(ws: Worksheet, headers: list[str], rows: list[list[Any]], start_row: int = 1,
                 widths: list[int] | None = None, wrap_cols: set[int] | None = None) -> int:
    ws.append([]) if start_row > ws.max_row + 1 else None
    for c, h in enumerate(headers, start=1):
        ws.cell(row=start_row, column=c, value=h)
    _style_header(ws, start_row, len(headers))
    r = start_row
    for r_off, row in enumerate(rows, start=1):
        r = start_row + r_off
        for c, value in enumerate(row, start=1):
            cell = _put(ws, r, c, value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=bool(wrap_cols and c in wrap_cols))
    if widths:
        for c, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(c)].width = w
    return r


def _status_formatting(ws: Worksheet, col_letter: str, first_row: int, last_row: int) -> None:
    rng = f"{col_letter}{first_row}:{col_letter}{last_row}"
    for value, fill, font in (
        (RecordStatus.VERIFIED.value, FILL_GREEN, FONT_GREEN),
        (RecordStatus.PARTIALLY_VERIFIED.value, FILL_BLUE, Font(color="1F4E78")),
        (RecordStatus.UNVERIFIED.value, FILL_GREY, FONT_GREY),
        (RecordStatus.NEEDS_REVIEW.value, FILL_RED, FONT_RED),
        (RecordStatus.EXCLUDED.value, FILL_YELLOW, FONT_YELLOW),
    ):
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{value}"'], fill=fill, font=font))


def _check_formatting(ws: Worksheet, col_letter: str, first_row: int, last_row: int) -> None:
    rng = f"{col_letter}{first_row}:{col_letter}{last_row}"
    anchor = f"{col_letter}{first_row}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({anchor},8)="verified"'], fill=FILL_GREEN, font=FONT_GREEN))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(LEFT({anchor},8)="conflict",LEFT({anchor},7)="invalid")'], fill=FILL_RED, font=FONT_RED))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({anchor},10)="unverified"'], fill=FILL_GREY, font=FONT_GREY))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{anchor}="missing"'], font=Font(color="A6A6A6", italic=True)))


def write_data_sheet(ws: Worksheet, records: list[Record], columns: list[Column], title: str) -> dict[str, str]:
    """Write the main sheet. Returns a map column-key -> column letter (for formulas)."""
    ws.freeze_panes = "C2"
    letters: dict[str, str] = {}
    for c, col in enumerate(columns, start=1):
        letter = get_column_letter(c)
        letters[col.key] = letter
        ws.cell(row=1, column=c, value=col.header)
        ws.column_dimensions[letter].width = col.width
    _style_header(ws, 1, len(columns))
    for r, rec in enumerate(records, start=2):
        for c, (col, value) in enumerate(zip(columns, record_row(rec, columns), strict=True), start=1):
            cell = _put(ws, r, c, value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=col.wrap)
            if col.number_format and isinstance(value, (int, float)):
                cell.number_format = col.number_format
            if col.hyperlink and isinstance(value, str) and value:
                target = f"mailto:{value}" if col.field is not None and col.field.type == "email" else value
                if target.startswith(("http://", "https://", "mailto:")):
                    cell.hyperlink = target
                    cell.font = Font(color="0563C1", underline="single")
    last = max(len(records) + 1, 2)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last}"
    _status_formatting(ws, letters["status"], 2, MAX_ROWS)
    if "completeness" in letters:
        col = letters["completeness"]
        ws.conditional_formatting.add(f"{col}2:{col}{MAX_ROWS}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color="5B9BD5"))
    for col in columns:
        if col.kind == "check":
            _check_formatting(ws, letters[col.key], 2, MAX_ROWS)
    if "reviewer_decision" in letters:
        dv = DataValidation(type="list", formula1='"Approved,Rejected,Needs fix,Contacted"', allow_blank=True,
                            showDropDown=False, promptTitle="Reviewer decision", prompt="Record the manual QA outcome for this row.")
        ws.add_data_validation(dv)
        col = letters["reviewer_decision"]
        dv.add(f"{col}2:{col}{MAX_ROWS}")
    for r in range(2, last + 1):
        ws.row_dimensions[r].height = None
    ws.sheet_properties.tabColor = "1F4E78"
    return letters


def write_review_sheet(ws: Worksheet, records: list[Record], columns: list[Column], schema: Schema) -> None:
    """Records that need a human decision, with the fields a reviewer needs and space for the outcome."""
    ws.sheet_properties.tabColor = "C00000"
    name_field = schema.name_field or "company_name"
    key_fields = [name_field] + [f.name for f in schema.fields if f.type in ("email", "phone", "url", "vat") and f.name != name_field]
    headers = ["Record ID", "Status", schema.field(name_field).display if schema.field(name_field) else "Name", "Review flags (why)"]
    headers += [schema.field(k).display for k in key_fields[1:] if schema.field(k)]
    headers += ["Checks"] + ["Source", "Source URL", "Resolution", "Resolved by", "Resolved on"]
    rows: list[list[Any]] = []
    for rec in records:
        checks = "\n".join(f"{k}: {rec.fields[k].status.value}{' - ' + rec.fields[k].note if rec.fields[k].note else ''}"
                           for k in key_fields if k in rec.fields and not rec.fields[k].is_empty)
        row = [rec.record_id, rec.status.value, rec.get(name_field), "\n".join(rec.flags)]
        row += [_display(rec, k, schema) for k in key_fields[1:] if schema.field(k)]
        row += [checks, rec.source, rec.source_url, None, None, None]
        rows.append(row)
    widths = [11, 18, 30, 60] + [22] * (len(key_fields) - 1) + [45, 12, 36, 24, 14, 12]
    last = _write_table(ws, headers, rows, widths=widths, wrap_cols={4, 5 + len(key_fields) - 1})
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(last, 2)}"
    _status_formatting(ws, "B", 2, MAX_ROWS)
    dv = DataValidation(type="list", formula1='"Fixed,Confirmed correct,Rejected - remove,Could not verify,Contacted company"', allow_blank=True)
    ws.add_data_validation(dv)
    res_col = get_column_letter(len(headers) - 2)
    dv.add(f"{res_col}2:{res_col}{MAX_ROWS}")
    if not rows:
        ws.cell(row=2, column=1, value="Nothing to review - every record passed the automatic checks.").font = SUBTLE_FONT


def write_duplicates_sheet(ws: Worksheet, groups: list[DuplicateGroup], by_id: dict[str, Record], schema: Schema) -> None:
    ws.sheet_properties.tabColor = "ED7D31"
    name_field = schema.name_field or "company_name"
    headers = ["Group", "Action", "Kept record", "Kept name", "Other record(s)", "Other name(s)", "Reason", "Name similarity %", "Decision"]
    rows = []
    for g in groups:
        master = by_id.get(g.master_id)
        others = [by_id[m] for m in g.member_ids if m != g.master_id and m in by_id]
        rows.append([
            g.group_id,
            "merged automatically" if g.merged else "POSSIBLE duplicate - decide",
            g.master_id,
            master.get(name_field) if master else "",
            ", ".join(o.record_id for o in others),
            "\n".join(str(o.get(name_field) or "") for o in others),
            g.reason,
            round(g.similarity),
            None,
        ])
    _write_table(ws, headers, rows, widths=[10, 26, 12, 30, 18, 30, 60, 12, 22], wrap_cols={6, 7})
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1='"Same entity - merge,Different entities - keep both,Unsure"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"I2:I{MAX_ROWS}")
    ws.conditional_formatting.add(f"B2:B{MAX_ROWS}", FormulaRule(formula=['LEFT(B2,8)="POSSIBLE"'], fill=FILL_YELLOW, font=FONT_YELLOW))
    if not rows:
        ws.cell(row=2, column=1, value="No duplicates were detected.").font = SUBTLE_FONT
    return None


def write_summary_sheet(ws: Worksheet, *, data_sheet: str, letters: dict[str, str], columns: list[Column],
                        report: RunReport, schema: Schema, project_title: str, records: list[Record]) -> None:
    ws.sheet_properties.tabColor = "70AD47"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 50
    ws["A1"] = project_title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Live summary - the counts below are formulas and update when the data sheet is edited."
    ws["A2"].font = SUBTLE_FONT
    q = f"'{data_sheet}'"
    id_col, st_col = letters["record_id"], letters["status"]
    total_ref = f"COUNTA({q}!{id_col}:{id_col})-1"
    r = 4
    ws.cell(row=r, column=1, value="Records").font = BOLD
    r += 1
    ws.cell(row=r, column=1, value="Delivered records")
    ws.cell(row=r, column=2, value=f"={total_ref}")
    total_cell = f"B{r}"
    r += 1
    for status in (RecordStatus.VERIFIED, RecordStatus.PARTIALLY_VERIFIED, RecordStatus.UNVERIFIED, RecordStatus.NEEDS_REVIEW):
        ws.cell(row=r, column=1, value=status.value.replace("_", " ").capitalize())
        ws.cell(row=r, column=2, value=f'=COUNTIF({q}!{st_col}:{st_col},"{status.value}")')
        ws.cell(row=r, column=3, value=f"=IF({total_cell}=0,0,B{r}/{total_cell})").number_format = "0.0%"
        r += 1
    ws.cell(row=r, column=1, value="Extracted before merge/exclusion")
    ws.cell(row=r, column=2, value=report.records_total)
    r += 1
    ws.cell(row=r, column=1, value="Duplicate groups")
    ws.cell(row=r, column=2, value=report.duplicate_groups)
    r += 1
    ws.cell(row=r, column=1, value="Reviewer decisions recorded")
    rd = letters.get("reviewer_decision")
    if rd:
        ws.cell(row=r, column=2, value=f"=COUNTA({q}!{rd}:{rd})-1")
    r += 2

    ws.cell(row=r, column=1, value="Field completeness").font = BOLD
    ws.cell(row=r, column=2, value="Filled").font = BOLD
    ws.cell(row=r, column=3, value="Share").font = BOLD
    ws.cell(row=r, column=4, value="How it is verified").font = BOLD
    r += 1
    for f in schema.fields:
        col = letters.get(f.name)
        if not col:
            continue
        ws.cell(row=r, column=1, value=f.display)
        ws.cell(row=r, column=2, value=f"=COUNTA({q}!{col}:{col})-1")
        ws.cell(row=r, column=3, value=f"=IF({total_cell}=0,0,B{r}/{total_cell})").number_format = "0.0%"
        ws.cell(row=r, column=4, value=f.verify or "").alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    ws.conditional_formatting.add(f"C5:C{r}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="70AD47"))
    r += 1

    ws.cell(row=r, column=1, value="Verification of key fields").font = BOLD
    for c, h in enumerate(("verified", "unverified", "conflict / invalid"), start=2):
        ws.cell(row=r, column=c, value=h).font = BOLD
    r += 1
    for col in columns:
        if col.kind != "check":
            continue
        letter = letters[col.key]
        ws.cell(row=r, column=1, value=col.header)
        ws.cell(row=r, column=2, value=f'=COUNTIF({q}!{letter}:{letter},"verified*")')
        ws.cell(row=r, column=3, value=f'=COUNTIF({q}!{letter}:{letter},"unverified*")')
        ws.cell(row=r, column=4, value=f'=COUNTIF({q}!{letter}:{letter},"conflict*")+COUNTIF({q}!{letter}:{letter},"invalid*")')
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="Sources").font = BOLD
    r += 1
    src_col = letters["source"]
    for source in sorted({rec.source for rec in records}):
        ws.cell(row=r, column=1, value=source)
        ws.cell(row=r, column=2, value=f'=COUNTIF({q}!{src_col}:{src_col},"{source}")')
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Generated").font = BOLD
    ws.cell(row=r, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M"))
    r += 1
    ws.cell(row=r, column=1, value="Status legend").font = BOLD
    r += 1
    legend = [
        ("VERIFIED", "Required fields present; contact details confirmed by an independent check or a second source.", FILL_GREEN),
        ("PARTIALLY VERIFIED", "At least one field confirmed; the rest is well-formed but single-source.", FILL_BLUE),
        ("UNVERIFIED", "Well-formed data from one source; nothing could be cross-checked.", FILL_GREY),
        ("NEEDS REVIEW", "A check failed, sources disagree, a required field is missing or a duplicate is suspected - see Review flags.", FILL_RED),
    ]
    for label, text, fill in legend:
        cell = ws.cell(row=r, column=1, value=label)
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
        ws.cell(row=r, column=2, value=text)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 30
        r += 1


def write_run_log_sheet(ws: Worksheet, report: RunReport, project_info: dict[str, Any]) -> None:
    ws.sheet_properties.tabColor = "7F7F7F"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 70
    ws["A1"] = "Run log"
    ws["A1"].font = TITLE_FONT
    r = 3
    d = report.to_dict()
    info = [
        ("Project", project_info.get("name")), ("Title", project_info.get("title")), ("Client / requested by", project_info.get("client")),
        ("Deadline", project_info.get("deadline")), ("Country rules", project_info.get("country")), ("Schema", d["settings"].get("schema")),
        ("Started (UTC)", report.started_at), ("Tool", d["tool"]), ("Python", d["python"]), ("Sources", ", ".join(d["settings"].get("sources", []))),
        ("Record limit", d["settings"].get("limit") or "none"), ("Records extracted", report.records_total), ("Records delivered", report.records_delivered),
    ]
    for k, v in info:
        ws.cell(row=r, column=1, value=k).font = BOLD
        _put(ws, r, 2, v if v not in (None, "") else "-")
        r += 1
    r += 1
    headers = ["Stage", "Input", "Output", "Seconds", "Notes"]
    rows = [[s["name"], s["input"], s["output"], s["seconds"], "; ".join(s["notes"])] for s in d["stages"]]
    r = _write_table(ws, headers, rows, start_row=r, wrap_cols={5}) + 2
    ws.cell(row=r, column=1, value="Warnings").font = BOLD
    r += 1
    if report.warnings:
        for w in report.warnings:
            _put(ws, r, 1, w)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
            ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True)
            r += 1
    else:
        ws.cell(row=r, column=1, value="none").font = SUBTLE_FONT
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Settings").font = BOLD
    r += 1
    for section in ("enrichment", "verification", "dedupe"):
        ws.cell(row=r, column=1, value=section)
        ws.cell(row=r, column=2, value=str(d["settings"].get(section)))
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True)
        r += 1


def write_dictionary_sheet(ws: Worksheet, columns: list[Column], schema: Schema) -> None:
    ws.sheet_properties.tabColor = "A5A5A5"
    headers = ["Column", "Field key", "Type", "Required", "Description", "How it is verified"]
    rows = []
    for col in columns:
        if col.field is not None and col.kind == "value":
            f = col.field
            rows.append([col.header, f.name, f.type, "yes" if f.required else "", f.description, f.verify])
        elif col.kind == "check":
            rows.append([col.header, col.key, "status", "", "Verification status of the field, with the reason and any other values seen in other sources.", ""])
        else:
            desc = {
                "record_id": "Stable identifier of the row (referenced in the Duplicates and Needs Review sheets).",
                "status": "Overall record status: VERIFIED, PARTIALLY_VERIFIED, UNVERIFIED or NEEDS_REVIEW (see Summary legend).",
                "flags": "Everything a human should look at, one issue per line. Lines starting with 'note:' are informational.",
                "completeness": "Share of schema fields that have a value (required fields weigh double).",
                "source": "Primary source the record was extracted from.",
                "source_url": "Link to the record in the primary source, for traceability.",
                "duplicate_group": "Set when the record was part of a duplicate cluster (see Duplicates sheet).",
                "extracted_at": "UTC timestamp of extraction.",
                "reviewer_decision": "Manual QA outcome (drop-down). Left empty by the tool.",
                "reviewer_notes": "Free text for the reviewer.",
            }.get(col.key, "")
            rows.append([col.header, col.key, "system", "", desc, ""])
    _write_table(ws, headers, rows, widths=[26, 20, 10, 9, 70, 60], wrap_cols={5, 6})
    ws.freeze_panes = "A2"


def export_excel(path: Path, records: list[Record], groups: list[DuplicateGroup], report: RunReport, schema: Schema,
                 project_info: dict[str, Any], *, include_excluded: bool = True) -> Path:
    columns = build_columns(schema)
    delivered = [r for r in records if r.status != RecordStatus.EXCLUDED]
    review = [r for r in records if r.status == RecordStatus.NEEDS_REVIEW or (include_excluded and r.status == RecordStatus.EXCLUDED and not r.duplicate_of)]
    by_id = {r.record_id: r for r in records}

    wb = Workbook()
    ws_data = wb.active
    ws_data.title = schema.label[:31]
    letters = write_data_sheet(ws_data, delivered, columns, schema.label)
    write_review_sheet(wb.create_sheet("Needs Review"), review, columns, schema)
    write_duplicates_sheet(wb.create_sheet("Duplicates"), groups, by_id, schema)
    write_summary_sheet(wb.create_sheet("Summary", 0), data_sheet=ws_data.title, letters=letters, columns=columns, report=report,
                        schema=schema, project_title=project_info.get("title") or schema.label, records=delivered)
    write_run_log_sheet(wb.create_sheet("Run Log"), report, project_info)
    write_dictionary_sheet(wb.create_sheet("Data Dictionary"), columns, schema)
    wb.active = 1  # open on the data sheet
    for ws in wb.worksheets:
        ws.sheet_view.selection[0].activeCell = "A1"
        ws.sheet_view.selection[0].sqref = "A1"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
