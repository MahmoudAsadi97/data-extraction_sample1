"""Import a seed list from CSV / Excel - e.g. a LinkedIn Sales Navigator or Apollo export,
a client-provided list, or names typed up during manual research.

The rows go through exactly the same normalisation, enrichment, verification
and duplicate detection as records extracted from the web.

Example::

    - type: csv_import
      path: data/input/seed_companies.csv
      mapping:                # CSV/Excel column -> schema field (optional; identical names map automatically)
        Company: company_name
        Web: website
        Tel: phone
      country: BE
"""

from __future__ import annotations

import csv
import logging
from collections.abc import Iterator
from pathlib import Path
from typing import Any

from ..models import RawRecord
from .base import BaseSource, SourceError, register

log = logging.getLogger(__name__)

# Common column names seen in LinkedIn / Apollo / CRM exports, mapped to schema fields.
AUTO_MAPPING = {
    "company": "company_name", "company name": "company_name", "organisation": "company_name",
    "organization": "company_name", "account name": "company_name", "name": "company_name",
    "website": "website", "web": "website", "url": "website", "domain": "website", "company website": "website",
    "email": "email", "e-mail": "email", "email address": "email",
    "phone": "phone", "telephone": "phone", "tel": "phone", "phone number": "phone", "company phone": "phone",
    "linkedin": "linkedin", "linkedin url": "linkedin", "company linkedin url": "linkedin", "linkedin_url": "linkedin",
    "facebook": "facebook", "facebook url": "facebook", "instagram": "instagram",
    "industry": "industry", "category": "category", "sector": "category",
    "city": "city", "town": "city", "postcode": "postcode", "postal code": "postcode", "zip": "postcode",
    "street": "street", "address": "street", "street address": "street", "country": "country",
    "vat": "vat_number", "vat number": "vat_number", "btw": "vat_number", "kbo": "vat_number",
    "enterprise number": "vat_number", "employees": "employees", "# employees": "employees",
    "founded": "founded", "founded year": "founded", "description": "description", "notes": "description",
}


MAX_INPUT_BYTES = 20 * 1024 * 1024
MAX_INPUT_ROWS = 5_000
MAX_INPUT_COLUMNS = 200


def _headers(values: list[Any]) -> list[str]:
    headers = [str(v).strip() if v is not None else "" for v in values]
    if not headers or len(headers) > MAX_INPUT_COLUMNS or any(not h for h in headers):
        raise ValueError("Use 1–200 non-empty column headers.")
    folded = [h.casefold() for h in headers]
    if len(set(folded)) != len(headers):
        raise ValueError("Duplicate column headers would overwrite data; give every column a unique name.")
    return headers


def read_rows(path: Path, encoding: str = "utf-8-sig", sheet: str | None = None) -> list[dict[str, Any]]:
    """Read bounded CSV/TSV/XLSX input without silently dropping malformed columns."""
    if path.stat().st_size > MAX_INPUT_BYTES:
        raise ValueError("Input exceeds the 20 MiB limit.")
    suffix = path.suffix.lower()
    rows: list[dict[str, Any]] = []

    def append(header: list[str], row: list[Any], line: int) -> None:
        if all(v in (None, "") for v in row):
            return
        if len(row) != len(header):
            raise ValueError(f"Row {line} has {len(row)} cells; expected {len(header)}.")
        if len(rows) >= MAX_INPUT_ROWS:
            raise ValueError("Input exceeds 5,000 rows; split it into smaller batches.")
        rows.append(dict(zip(header, row, strict=True)))

    if suffix in (".xlsx", ".xlsm"):
        from zipfile import BadZipFile, ZipFile

        from openpyxl import load_workbook

        try:
            with ZipFile(path) as archive:
                if sum(f.file_size for f in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError("Workbook expands beyond the 100 MiB limit.")
            wb = load_workbook(path, read_only=True, data_only=False)
        except BadZipFile as exc:
            raise ValueError("Invalid XLSX workbook.") from exc
        try:
            if sheet and sheet not in wb.sheetnames:
                raise ValueError(f"Sheet does not exist: {sheet}")
            ws = wb[sheet] if sheet else wb.active
            if ws.max_column and ws.max_column > MAX_INPUT_COLUMNS:
                raise ValueError("Workbook exceeds 200 columns.")
            if ws.max_row and ws.max_row > MAX_INPUT_ROWS + 1:
                raise ValueError("Workbook exceeds 5,000 rows; remove unused formatted rows or split the input.")
            rows_iter = ws.iter_rows()
            first = next(rows_iter, [])
            if any(c.data_type == "f" for c in first):
                raise ValueError("Column headers cannot contain formulas.")
            header = _headers([c.value for c in first])
            for line, cells in enumerate(rows_iter, 2):
                if any(c.data_type == "f" for c in cells):
                    raise ValueError(f"Row {line} contains a formula; paste values before importing.")
                append(header, [c.value for c in cells], line)
            return rows
        finally:
            wb.close()
    if suffix not in (".csv", ".tsv"):
        raise ValueError("Supported inputs: .csv, .tsv, .xlsx, .xlsm.")
    with path.open("r", encoding=encoding, newline="") as fh:
        sample = fh.read(4096)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
        reader = csv.reader(fh, dialect=dialect, strict=True)
        header = _headers(next(reader, []))
        try:
            for row in reader:
                append(header, row, reader.line_num)
        except csv.Error as exc:
            raise ValueError(f"Malformed CSV at line {reader.line_num}: {exc}") from exc
        return rows


@register
class CsvImportSource(BaseSource):
    type = "csv_import"
    description = "Seed list from CSV/Excel (client list, LinkedIn/Apollo export, manual research) - enriched & verified like any other source."
    produces = "leads"

    def extract(self, limit: int | None = None) -> Iterator[RawRecord]:
        raw_path = self.require_opt("path")
        path = self.project.resolve_path(raw_path)
        if not path.exists():
            raise SourceError(f"csv_import: file not found: {path}")
        mapping: dict[str, str] = {str(k).strip().lower(): v for k, v in (self.opt("mapping") or {}).items()}
        rows = read_rows(path, encoding=self.opt("encoding", "utf-8-sig"), sheet=self.opt("sheet"))
        log.info("csv_import: %d rows read from %s", len(rows), path.name)
        country = self.opt("country") or self.project.project.country
        from ..qa import map_columns

        column_map, ignored = map_columns(list(rows[0]) if rows else [], self.project.schema_def, mapping)
        unmapped = set(ignored)
        for i, row in enumerate(rows, start=2):
            values = {column_map[column]: value for column, value in row.items()
                      if column in column_map and value not in (None, "")}
            if not values:
                continue
            values.setdefault("country", country)
            yield RawRecord(
                source=self.label,
                source_id=f"row {i}",
                source_url=path.name,
                values=values,
                raw={"row": i, "columns": row},
            )
            if limit and (i - 1) >= limit:
                break
        if unmapped:
            self.warn("columns ignored (no schema field): " + ", ".join(sorted(unmapped)))
